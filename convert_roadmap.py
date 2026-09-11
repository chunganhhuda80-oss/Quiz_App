#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
Chuyển đổi file ROADMAPNETWORK_AI.xlsx sang roadmap_data.json
Phục vụ hiển thị sơ đồ tư duy (Mindmap) và chi tiết kiến thức từng nhánh theo tuần.
"""

import os
import sys
import re
import json
import openpyxl

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

def parse_week_content(week_id, week_title_raw, raw_text):
    if not raw_text or not str(raw_text).strip():
        return {
            'weekId': week_id,
            'hasData': False,
            'title': week_title_raw,
            'fullTitle': week_title_raw,
            'domain': 'Networking' if week_id <= 8 else 'ML/DL',
            'groups': [],
            'lab': None
        }
    
    domain = 'Networking' if week_id <= 8 else 'ML/DL'
    clean_title = re.sub(r'^Tuần\s*\d+\s*[:\-]\s*', '', str(week_title_raw)).strip()
    
    lines = str(raw_text).split('\n')
    groups = []
    lab_tasks = []
    
    current_group = None
    current_item = None
    in_lab = False
    
    for line in lines:
        stripped = line.strip()
        if not stripped:
            continue
            
        # Kiểm tra nếu là phần Lab
        if re.match(r'^Lab\s+tuần\s+này', stripped, re.IGNORECASE):
            in_lab = True
            current_group = None
            current_item = None
            continue
            
        if in_lab:
            if stripped.startswith('- ') or stripped.startswith('* '):
                lab_tasks.append(stripped[2:].strip())
            else:
                lab_tasks.append(stripped)
            continue
            
        # Kiểm tra dòng Nhóm X:
        group_match = re.match(r'^(Nhóm\s*\d+)\s*[:\-]\s*(.*)$', stripped, re.IGNORECASE)
        if group_match:
            current_group = {
                'badge': group_match.group(1).strip(),
                'name': group_match.group(2).strip(),
                'items': []
            }
            groups.append(current_group)
            current_item = None
            continue
            
        # Kiểm tra dòng Mục con: 1. Tiêu đề hoặc a. Tiêu đề
        item_match = re.match(r'^(\d+\.|\b[a-zA-Z]\.)\s+(.*)$', stripped)
        if item_match and current_group is not None:
            current_item = {
                'index': item_match.group(1).strip(),
                'title': item_match.group(2).strip(),
                'details': []
            }
            current_group['items'].append(current_item)
            continue
            
        # Các dòng chi tiết gạch đầu dòng hoặc mô tả
        if current_item is not None:
            current_item['details'].append(stripped)
        elif current_group is not None:
            if not current_group['items']:
                current_group['items'].append({
                    'index': '1.',
                    'title': 'Tổng quan & Kiến thức cốt lõi',
                    'details': [stripped]
                })
            else:
                current_group['items'][-1]['details'].append(stripped)
                
    return {
        'weekId': week_id,
        'hasData': True,
        'title': clean_title or week_title_raw,
        'fullTitle': week_title_raw,
        'domain': domain,
        'groups': groups,
        'lab': lab_tasks if lab_tasks else None
    }

def main():
    excel_path = 'ROADMAPNETWORK_AI.xlsx'
    if not os.path.exists(excel_path):
        print(f"Error: {excel_path} not found")
        sys.exit(1)

    wb = openpyxl.load_workbook(excel_path, data_only=True)
    sheet = wb['RoadMap']

    parsed_weeks = []
    for c in range(1, 16):
        title_cell = sheet.cell(3, c).value or f'Tuần {c}'
        content_cell = sheet.cell(4, c).value
        parsed = parse_week_content(c, title_cell, content_cell)
        parsed_weeks.append(parsed)
        total_items = sum(len(g['items']) for g in parsed['groups'])
        lab_count = len(parsed['lab']) if parsed['lab'] else 0
        print(f"Week {c:2d} | Domain: {parsed['domain']:10s} | HasData: {str(parsed['hasData']):5s} | Groups: {len(parsed['groups'])} | Items: {total_items:2d} | Labs: {lab_count}")

    output_path = 'roadmap_data.json'
    with open(output_path, 'w', encoding='utf-8') as f:
        json.dump(parsed_weeks, f, ensure_ascii=False, indent=2)

    print(f"\n[OK] Đã chuyển đổi thành công sang {output_path}")

if __name__ == '__main__':
    main()
